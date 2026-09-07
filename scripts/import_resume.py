# /// script
# requires-python = ">=3.10"
# dependencies = ["pypdf", "python-docx", "striprtf"]
# ///
"""把使用者丟過來的履歷轉成 Markdown，存進 ~/.nextrole/resumes/master-<id>.md。

    uv run import_resume.py ~/Downloads/我的履歷.pdf --id pm --label "PM 版" --for "產品規劃、需求管理"
    pbpaste | uv run import_resume.py --stdin --id pm

支援 .md/.markdown/.txt（直接收）、.pdf、.docx、.rtf、.html/.htm。
⚠️ 轉出來的東西一定要給使用者看過再用 —— PDF 抽文字常常掉排版、把兩欄併成一行。
"""
from __future__ import annotations

import argparse
import html as _html
import os
import re
import sys

import store

TEXTLIKE = {".md", ".markdown", ".txt", ".text"}


def _from_pdf(path: str) -> str:
    from pypdf import PdfReader
    pages = [(p.extract_text() or "") for p in PdfReader(path).pages]
    return "\n\n".join(pages)


def _from_docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    out = []
    for p in d.paragraphs:
        t = p.text.rstrip()
        style = (p.style.name or "").lower()
        if t and style == "title":          # Word 的「標題」樣式不叫 heading
            out.append("# " + t)
        elif t and style.startswith("heading"):
            lvl = "".join(c for c in style if c.isdigit()) or "2"
            out.append("#" * min(6, int(lvl) + 1) + " " + t)
        elif t and style.startswith("list"):
            out.append("- " + t)
        else:
            out.append(t)
    for tbl in d.tables:                      # 表格轉成條列，不硬做成 Markdown 表格
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append("- " + "｜".join(cells))
    return "\n".join(out)


def _from_rtf(path: str) -> str:
    from striprtf.striprtf import rtf_to_text
    with open(path, encoding="utf-8", errors="ignore") as f:
        return rtf_to_text(f.read())


def _from_html(path: str) -> str:
    with open(path, encoding="utf-8", errors="ignore") as f:
        raw = f.read()
    raw = re.sub(r"(?is)<(script|style).*?</\1>", "", raw)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</(p|div|li|tr|h[1-6])>", "\n", raw)
    raw = re.sub(r"(?i)<li[^>]*>", "- ", raw)
    return _html.unescape(re.sub(r"<[^>]+>", "", raw))


READERS = {".pdf": _from_pdf, ".docx": _from_docx, ".rtf": _from_rtf,
           ".html": _from_html, ".htm": _from_html}

# 轉不了的，各自給一條使用者做得到的出路，不要只說「不支援」
DEAD_ENDS = {
    ".doc": "舊版 Word。用 Word 或 Pages 另存成 .docx 再丟一次，"
            "或在終端機跑 `textutil -convert docx 檔名.doc`。",
    ".pages": "Pages 檔。在 Pages 裡「輸出」成 PDF 或 Word 再丟一次。",
    ".png": "圖片沒辦法抽文字，這個工具不做 OCR。",
    ".jpg": "圖片沒辦法抽文字，這個工具不做 OCR。",
    ".jpeg": "圖片沒辦法抽文字，這個工具不做 OCR。",
}


def tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    lines = [ln.rstrip() for ln in text.split("\n")]
    return "\n".join(lines).strip() + "\n"


def extract(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXTLIKE:
        with open(path, encoding="utf-8", errors="replace") as f:
            return f.read()
    if ext in DEAD_ENDS:
        raise SystemExit(f"轉不了 {ext}：{DEAD_ENDS[ext]}")
    reader = READERS.get(ext)
    if reader is None:
        raise SystemExit(
            f"不知道怎麼讀 {ext or '（沒有副檔名）'}。"
            "另存成 PDF、Word（.docx）或純文字再丟一次，或直接把內容貼給我。"
        )
    return reader(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path", nargs="?", help="履歷檔案路徑")
    ap.add_argument("--stdin", action="store_true", help="改從 stdin 讀（貼上的純文字）")
    ap.add_argument("--id", default="main", help="版本代號，會變成檔名 master-<id>.md")
    ap.add_argument("--label", default=None, help="版本名稱，例如「PM 版」")
    ap.add_argument("--for", dest="for_", default=None, help="這一版是投什麼方向的")
    ap.add_argument("--force", action="store_true", help="已經有同名檔案時覆寫")
    args = ap.parse_args()

    if not args.stdin and not args.path:
        raise SystemExit("要給檔案路徑，或加 --stdin 從標準輸入讀。")

    raw = sys.stdin.read() if args.stdin else extract(os.path.expanduser(args.path))
    text = tidy(raw)
    if len(text.strip()) < 50:
        raise SystemExit(
            "抽出來的內容少於 50 字，多半是掃描檔或純圖片的 PDF。"
            "這個工具不做 OCR —— 請改丟文字版，或直接把內容貼給我。"
        )

    store.ensure_dirs()
    out = os.path.join(store.ROOT, "resumes", f"master-{args.id}.md")
    if os.path.exists(out) and not args.force:
        raise SystemExit(f"{out} 已經存在。要覆寫請加 --force，或換一個 --id。")
    with open(out, "w", encoding="utf-8") as f:
        f.write(text)

    if args.label:
        cfg = store.load_config()
        vers = [v for v in cfg.get("resume_versions", []) if v.get("id") != args.id]
        vers.append({"id": args.id, "label": args.label, "for": args.for_ or ""})
        cfg["resume_versions"] = vers
        store.save_config(cfg)

    src = "貼上的文字" if args.stdin else args.path
    print(f"來源：{src}")
    print(f"寫出：{out}（{len(text.strip())} 字）")
    if args.label:
        print(f"已登記版本：{args.id} — {args.label}")

    # 警告要對得上實際風險，不然使用者會學會忽略它
    ext = "" if args.stdin else os.path.splitext(args.path)[1].lower()
    if ext == ".pdf":
        print("\n⚠️ 把轉出來的內容給使用者看過再用。PDF 抽文字常常掉排版、"
              "把兩欄併成一行、或漏掉表格裡的字。")
    elif ext in (".docx", ".rtf", ".html", ".htm"):
        print("\n⚠️ 把轉出來的內容給使用者看過再用。表格會被拆成條列，"
              "版面順序可能跟原檔不一樣。")


if __name__ == "__main__":
    main()
